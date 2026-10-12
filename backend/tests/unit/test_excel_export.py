from io import BytesIO

import httpx
import pytest
from fastapi import FastAPI
from openpyxl import load_workbook
from pydantic import ValidationError

from app.auth.dependencies import get_current_organization
from app.core.errors import install_exception_handlers
from app.exports.excel import ExcelTableExport, build_excel, export_filename
from app.exports.router import EXCEL_MEDIA_TYPE, router


def test_excel_preserves_values_and_never_executes_text_as_formulas():
    table = ExcelTableExport(
        title="Backlog / Japan [current]",
        headers=["Customer", "Amount", "Active", "Missing", "Order ID"],
        rows=[
            ['=HYPERLINK("https://example.com")', 1234.5678, True, None, 123456789012345678],
            ["東京 & Singapore", -42, False, "00123", "=SUM(A1:A2)"],
        ],
    )
    book = load_workbook(BytesIO(build_excel(table)))
    sheet = book.active
    assert list(sheet.values) == [
        tuple(table.headers),
        ('=HYPERLINK("https://example.com")', 1234.5678, True, None, "123456789012345678"),
        ("東京 & Singapore", -42, False, "00123", "=SUM(A1:A2)"),
    ]
    assert sheet["A2"].data_type == sheet["E3"].data_type == "s"
    assert sheet["B2"].data_type == "n"
    assert sheet["C2"].data_type == "b"
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:E3"
    assert len(sheet.title) <= 31
    assert "/" not in sheet.title


def test_export_includes_rows_beyond_a_preview_and_marks_truncation():
    table = ExcelTableExport(
        title="Orders", headers=["Order"], rows=[[i] for i in range(600)], truncated=True
    )
    sheet = load_workbook(BytesIO(build_excel(table))).active
    assert sheet["A601"].value == 599
    assert "truncated" in sheet["A603"].value
    assert sheet.auto_filter.ref == "A1:A601"


def test_empty_table_exports_headers():
    sheet = load_workbook(BytesIO(build_excel(ExcelTableExport(headers=["ID"], rows=[])))).active
    assert list(sheet.values) == [("ID",)]


@pytest.mark.parametrize("rows", [[[1, 2]], [[float("inf")]], [["x" * 32768]]])
def test_invalid_cells_or_dimensions_are_rejected(rows):
    with pytest.raises(ValidationError):
        ExcelTableExport(headers=["ID"], rows=rows)


def test_filename_is_safe():
    assert export_filename('Report\r\n/"invalid"') == "Report_invalid.xlsx"
    assert export_filename("...") == "Table.xlsx"


async def test_authenticated_export_endpoint_returns_an_excel_download():
    app = FastAPI()
    install_exception_handlers(app)
    app.include_router(router, prefix="/api/v1")

    async def organization():
        return None

    app.dependency_overrides[get_current_organization] = organization
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        result = await client.post(
            "/api/v1/exports/excel", json={"title": "Orders", "headers": ["ID"], "rows": [[7]]}
        )
        assert result.status_code == 200
        assert result.headers["content-type"] == EXCEL_MEDIA_TYPE
        assert result.headers["content-disposition"] == 'attachment; filename="Orders.xlsx"'
        assert result.headers["cache-control"] == "no-store"
        assert load_workbook(BytesIO(result.content)).active["A2"].value == 7
        invalid = await client.post(
            "/api/v1/exports/excel", json={"headers": ["ID"], "rows": [[1, 2]]}
        )
        assert invalid.status_code == 422


async def test_export_requires_authentication():
    app = FastAPI()
    install_exception_handlers(app)
    app.include_router(router, prefix="/api/v1")
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        result = await client.post("/api/v1/exports/excel", json={"headers": ["ID"], "rows": [[7]]})
        assert result.status_code == 401
