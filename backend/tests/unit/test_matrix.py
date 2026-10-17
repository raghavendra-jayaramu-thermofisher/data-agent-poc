from __future__ import annotations

import json
import uuid
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

import duckdb
import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from app.agent.matrix import (
    CreateMatrixArgs,
    MatrixError,
    MatrixViewState,
    build_matrix,
    matrix_projection,
)
from app.exports.matrix import build_matrix_excel


def example():
    spec = CreateMatrixArgs.model_validate(
        {
            "sql_execution_id": "11111111-1111-4111-8111-111111111111",
            "title": "Orders by period",
            "rows": [{"field": "division"}, {"field": "category"}],
            "columns": [{"field": "period", "sort_field": "period_order", "sort_type": "number"}],
            "measures": [{"field": "amount", "decimal_places": 2}, {"field": "units"}],
            "row_depth_field": "rd",
            "column_depth_field": "cd",
        }
    )
    columns = ["division", "category", "period", "period_order", "amount", "units", "rd", "cd"]
    rows = [
        ["North", "Parts", "Week 10", 10, "9007199254740993.12", 2, 2, 1],
        ["North", None, "Week 2", 2, "0.10", 0, 2, 1],
        ["South", "Parts", "Week 2", 2, None, 3, 2, 1],
        ["North", None, "Week 10", 10, "9007199254740993.12", 2, 1, 1],
        ["North", None, "Week 2", 2, "0.10", 0, 1, 1],
        ["South", None, "Week 2", 2, None, 3, 1, 1],
        [None, None, "Week 10", 10, "9007199254740993.12", 2, 0, 1],
        [None, None, "Week 2", 2, "0.10", 3, 0, 1],
    ]
    return spec, columns, rows


def test_exact_values_null_categories_numeric_order_and_sql_totals():
    spec, columns, rows = example()
    payload = build_matrix(spec, columns, rows)
    assert [node["label"] for node in payload["column_nodes"]] == ["Week 2", "Week 10"]
    assert any(node["label"] == "(Blank)" and node["depth"] == 2 for node in payload["row_nodes"])
    assert "9007199254740993.12" in json.dumps(payload)
    assert payload["totals_source"] == "sql"
    assert len(payload["cells"]) == len(rows)
    assert build_matrix(spec, columns, list(reversed(rows))) == {
        **payload,
        "source_digest": build_matrix(spec, columns, list(reversed(rows)))["source_digest"],
    }


@pytest.mark.parametrize(
    "case,code",
    [
        ("duplicate", "MATRIX_DUPLICATE_CELL"),
        ("missing_total", "MATRIX_TOTALS_MISSING"),
        ("invalid_number", "MATRIX_VALUE_INVALID"),
        ("boolean", "MATRIX_VALUE_INVALID"),
        ("bad_depth", "MATRIX_DEPTH_INVALID"),
        ("bad_width", "MATRIX_SOURCE_INVALID"),
        ("sort_conflict", "MATRIX_SORT_CONFLICT"),
    ],
)
def test_rejects_ambiguous_or_incomplete_sources(case, code):
    spec, columns, rows = example()
    if case == "duplicate":
        rows.append(rows[0])
    if case == "missing_total":
        rows.pop()
    if case == "invalid_number":
        rows[0][4] = "NaN"
    if case == "boolean":
        rows[0][4] = True
    if case == "bad_depth":
        rows[0][-2] = 4
    if case == "bad_width":
        rows[0].pop()
    if case == "sort_conflict":
        rows[2][3] = 3
    with pytest.raises(MatrixError) as error:
        build_matrix(spec, columns, rows)
    assert error.value.code == code


def test_schema_requires_totals_contract_and_disjoint_fields():
    spec, _, _ = example()
    for changes in ({"row_depth_field": None}, {"measures": [{"field": "period"}]}):
        with pytest.raises(ValidationError):
            CreateMatrixArgs.model_validate({**spec.model_dump(), **changes})


def test_truncation_and_size_limits_fail_closed():
    spec, columns, rows = example()
    for kwargs in (
        {"truncated": True},
        {"max_columns": 1},
        {"max_cells": 2},
        {"max_bytes": 100},
        {"max_nodes": 1},
    ):
        with pytest.raises(MatrixError):
            build_matrix(spec, columns, rows, **kwargs)


def test_field_mapping_cannot_supply_numbers_or_nonexistent_fields():
    spec, columns, rows = example()
    with pytest.raises(ValidationError):
        CreateMatrixArgs.model_validate({**spec.model_dump(), "data": [["invented", 100]]})
    with pytest.raises(MatrixError, match="exact output fields"):
        build_matrix(spec, ["missing", *columns[1:]], rows)


def test_leaf_only_empty_matrix_and_literal_null_label():
    spec = CreateMatrixArgs.model_validate(
        {
            "sql_execution_id": uuid.uuid4(),
            "title": "Measurements",
            "rows": [{"field": "sensor"}],
            "columns": [{"field": "date", "sort_type": "date"}],
            "measures": [{"field": "reading"}],
            "show_row_grand_total": False,
            "show_row_subtotals": False,
        }
    )
    assert build_matrix(spec, ["sensor", "date", "reading"], [])["cells"] == []
    payload = build_matrix(
        spec,
        ["sensor", "date", "reading"],
        [
            [None, "2026-01-01", 0],
            ["(Blank)", "2026-01-01", -2],
        ],
    )
    assert len({node["id"] for node in payload["row_nodes"]}) == 2


def test_nonadditive_total_is_taken_from_sql_without_summing_children():
    spec, columns, rows = example()
    rows[-1][4] = "0.075"  # A database-calculated measure, not a sum of child cells.
    payload = build_matrix(spec, columns, rows)
    total = next(node["id"] for node in payload["row_nodes"] if node["depth"] == 0)
    assert any(
        cell["row_id"] == total and cell["values"][0] == "0.075" for cell in payload["cells"]
    )


def test_expansion_export_projection_and_formula_safety():
    spec, columns, rows = example()
    rows[0][1] = '=HYPERLINK("https://example.invalid")'
    payload = build_matrix(spec, columns, rows)
    north = next(node for node in payload["row_nodes"] if node["path"] == ["North"])
    state = MatrixViewState(collapsed_rows=[north["id"]])
    visible, _ = matrix_projection(payload, state)
    assert not any(node["path"][:1] == ["North"] and node["depth"] == 2 for node in visible)
    workbook = load_workbook(BytesIO(build_matrix_excel(payload, state)))
    sheet = workbook.active
    assert sheet.freeze_panes == "B3"
    assert sheet.max_row == 2 + len(visible)
    full = load_workbook(BytesIO(build_matrix_excel(payload, MatrixViewState()))).active
    assert any(cell.value == rows[0][1] and cell.data_type == "s" for row in full for cell in row)
    assert any(
        cell.value == "9,007,199,254,740,993.12"
        and cell.data_type == "s"
        and cell.comment.text == "Exact SQL value: 9007199254740993.12"
        for row in full
        for cell in row
    )
    with pytest.raises(MatrixError):
        matrix_projection(payload, MatrixViewState(collapsed_rows=["r" + "0" * 24]))


def sensor_example():
    # Unrelated data/table: average handling comes from the SQL engine, never application rules.
    with duckdb.connect() as db:
        result = db.execute("""
            WITH observations(team, device, year, month, reading) AS (VALUES
                ('A', 'x', 2026, 1, 10.0), ('A', 'x', 2026, 1, 20.0),
                ('A', 'y', 2026, 1, 100.0), ('B', 'z', 2026, 2, 5.0))
            SELECT team, device, year, month, AVG(reading) AS mean_reading,
                2-GROUPING(team)-GROUPING(device) AS rd,
                2-GROUPING(year)-GROUPING(month) AS cd
            FROM observations GROUP BY GROUPING SETS (
                (team,device,year,month), (team,year,month), (year,month),
                (team,device,year), (team,year), (year),
                (team,device), (team), ())
        """)
        columns = [column[0] for column in result.description]
        rows = [list(row) for row in result.fetchall()]
    spec = CreateMatrixArgs.model_validate(
        {
            "sql_execution_id": "33333333-3333-4333-8333-333333333333",
            "title": "Sensor averages",
            "rows": [{"field": "team"}, {"field": "device"}],
            "columns": [
                {"field": "year", "sort_type": "number"},
                {"field": "month", "sort_type": "number"},
            ],
            "measures": [{"field": "mean_reading", "decimal_places": 2}],
            "row_depth_field": "rd",
            "column_depth_field": "cd",
            "show_column_subtotals": True,
            "show_column_grand_total": True,
        }
    )
    rows.sort(key=json.dumps)
    return spec, columns, rows


def test_native_grouping_sets_average_and_multilevel_columns():
    spec, columns, rows = sensor_example()
    payload = build_matrix(spec, columns, rows)
    grand_row = next(node["id"] for node in payload["row_nodes"] if node["depth"] == 0)
    grand_col = next(node["id"] for node in payload["column_nodes"] if node["depth"] == 0)
    assert next(
        cell["values"]
        for cell in payload["cells"]
        if (cell["row_id"], cell["column_id"]) == (grand_row, grand_col)
    ) == ["33.75"]
    year = next(node for node in payload["column_nodes"] if node["depth"] == 1)
    _, visible = matrix_projection(payload, MatrixViewState(collapsed_columns=[year["id"]]))
    assert len(visible) == 2  # Year subtotal and grand total, not the two months.


@pytest.mark.asyncio
async def test_tool_authorization_and_persistence(monkeypatch):
    from app.agent import tools
    from app.core.config import Settings

    spec, columns, rows = example()
    source_id = uuid.uuid4()
    source = SimpleNamespace(
        id=spec.sql_execution_id,
        status="succeeded",
        data_source_id=source_id,
        columns=columns,
        row_count=len(rows),
        truncated=False,
    )
    table = SimpleNamespace(
        payload={
            "schema_version": 1,
            "columns": columns,
            "rows": rows,
            "row_count": len(rows),
            "data_source_id": str(source_id),
            "truncated": False,
            "model_preview_truncated": True,
            "display": "auto",
        }
    )
    persisted = []

    def add(item):
        item.id = uuid.uuid4()
        persisted.append(item)

    repo = SimpleNamespace(
        get_sql_for_job=AsyncMock(return_value=source),
        get_table_artifact_for_sql=AsyncMock(return_value=table),
        list_artifacts=AsyncMock(return_value=persisted),
        add_artifact=add,
    )
    session = SimpleNamespace(commit=AsyncMock())

    class SessionContext:
        async def __aenter__(self):
            return session

        async def __aexit__(self, *_):
            pass

    monkeypatch.setattr(tools, "get_session_factory", lambda: SessionContext)
    monkeypatch.setattr(tools, "AgentExecutionRepository", lambda _: repo)
    monkeypatch.setattr(tools, "append_agent_event", AsyncMock())
    context = tools.AgentToolContext(
        uuid.uuid4(), uuid.uuid4(), uuid.uuid4(), frozenset(), Settings()
    )
    with pytest.raises(tools.AgentToolError) as error:
        await tools.create_matrix(context, spec.model_dump(mode="json"), uuid.uuid4())
    assert error.value.code == "MATRIX_SOURCE_NOT_FOUND"
    from dataclasses import replace

    context = replace(context, allowed_data_source_ids=frozenset([source_id]))
    response = await tools.create_matrix(context, spec.model_dump(mode="json"), uuid.uuid4())
    assert response["ok"] is True
    assert table.payload["display"] == "hidden"
    assert table.payload["model_preview_truncated"] is True
    again = await tools.create_matrix(context, spec.model_dump(mode="json"), uuid.uuid4())
    assert again["matrix_artifact_id"] == response["matrix_artifact_id"]
    assert len(persisted) == 1
    table.payload = {**table.payload, "truncated": True, "display": "auto"}
    source.truncated = True
    with pytest.raises(tools.AgentToolError) as error:
        await tools.create_matrix(context, spec.model_dump(mode="json"), uuid.uuid4())
    assert error.value.code == "MATRIX_SOURCE_TRUNCATED"
    assert table.payload["display"] == "auto"


@pytest.mark.asyncio
@pytest.mark.parametrize("tampered", [False, True])
async def test_matrix_is_verified_and_visible_to_the_final_answer(monkeypatch, tampered):
    from app.agent import evidence
    from app.agent.final_answer_validation import validate_final_answer
    from app.agent.models import JobArtifact, SQLExecution, ToolExecution

    spec, columns, rows = example()
    payload = build_matrix(spec, columns, rows)
    if tampered:
        payload["cells"][0]["values"][0] = "99999"
    job_id, source_id, tool_id = [uuid.uuid4() for _ in range(3)]
    source = SQLExecution(
        id=spec.sql_execution_id,
        tool_execution_id=tool_id,
        data_source_id=source_id,
        status="succeeded",
        columns=columns,
        row_count=len(rows),
        truncated=False,
        sql_text="SELECT SUM(amount) FROM measurements",
    )
    table = JobArtifact(
        id=uuid.uuid4(),
        job_id=job_id,
        kind="table",
        name="Source",
        source_sql_execution_id=source.id,
        payload={
            "schema_version": 1,
            "columns": columns,
            "rows": rows,
            "row_count": len(rows),
            "data_source_id": str(source_id),
            "truncated": False,
            "display": "hidden",
            "model_preview_truncated": True,
        },
    )
    artifact = JobArtifact(
        id=uuid.uuid4(),
        job_id=job_id,
        kind="matrix",
        name=spec.title,
        source_sql_execution_id=source.id,
        payload=payload,
    )
    tool = ToolExecution(
        id=tool_id,
        job_id=job_id,
        sequence=1,
        tool_name="execute_sql",
        status="succeeded",
        output_summary={"source_kind": "postgres"},
    )
    repo = SimpleNamespace(
        list_tools=AsyncMock(return_value=[tool]),
        list_sql=AsyncMock(return_value=[source]),
        list_artifacts=AsyncMock(return_value=[table, artifact]),
    )
    monkeypatch.setattr(evidence, "AgentExecutionRepository", lambda _: repo)
    bundle = await evidence.collect_evidence(
        object(), job_id=job_id, question="Compare measurements", analyst_notes=""
    )
    assert bundle.artifacts[1].valid is not tampered
    if not tampered:
        assert [item.kind for item in bundle.visible_artifacts] == ["matrix"]
        assert evidence.validate_evidence(bundle) == []
        assert validate_final_answer(f"## Results\n\n[[artifact:{artifact.id}]]", bundle) == []
    else:
        assert any("matrix" in issue for issue in evidence.validate_evidence(bundle))


def test_shared_frontend_fixture_matches_backend_contract():
    fixture = json.loads((Path(__file__).parents[3] / "test-fixtures/matrix.json").read_text())
    spec, columns, rows = example()
    assert fixture == build_matrix(spec, columns, rows)
