"""Export a stored matrix's visible hierarchy without querying business data."""

from decimal import ROUND_HALF_UP, Decimal, localcontext
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.agent.matrix import CreateMatrixArgs, MatrixViewState, matrix_projection


def build_matrix_excel(payload: dict, state: MatrixViewState) -> bytes:
    spec = CreateMatrixArgs.model_validate(payload["spec"])
    rows, columns = matrix_projection(payload, state)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Matrix"
    workbook.properties.title = spec.title
    workbook.properties.description = (
        "SQL-provided values and totals; current matrix expansion state."
    )
    header_depth = len(spec.columns) + (1 if len(spec.measures) > 1 else 0)
    sheet.freeze_panes = f"B{header_depth + 1}"
    sheet.column_dimensions["A"].width = 38
    cells = {(cell["row_id"], cell["column_id"]): cell["values"] for cell in payload["cells"]}

    def text(row: int, column: int, value: str):
        cell = sheet.cell(row, column, ILLEGAL_CHARACTERS_RE.sub("", value))
        cell.data_type = "s"  # SQL labels that start with '=' remain text, never formulas.
        return cell

    text(1, 1, " / ".join(item.label or item.field for item in spec.rows))
    if header_depth > 1:
        sheet.merge_cells(start_row=1, start_column=1, end_row=header_depth, end_column=1)
    headers: list[list[tuple[tuple, str]]] = [[] for _ in range(header_depth)]
    for node in columns:
        for measure in spec.measures:
            for level in range(len(spec.columns)):
                path = node["path"]
                label = (
                    "Grand total"
                    if not path
                    else ("(Blank)" if path[level] is None else str(path[level]))
                    if level < len(path)
                    else "Total"
                )
                # Prefix identity prevents merging repeated labels under different parents.
                identity = (
                    (node["id"], level)
                    if level >= len(path)
                    else tuple((type(value).__name__, str(value)) for value in path[: level + 1])
                )
                headers[level].append((identity, label))
            if len(spec.measures) > 1:
                headers[-1].append(((node["id"], measure.field), measure.label or measure.field))
    for level, header in enumerate(headers, start=1):
        index = 0
        while index < len(header):
            end = index + 1
            while end < len(header) and header[end][0] == header[index][0]:
                end += 1
            text(level, index + 2, header[index][1])
            if end > index + 1:
                sheet.merge_cells(
                    start_row=level, start_column=index + 2, end_row=level, end_column=end + 1
                )
            index = end
    for row in sheet.iter_rows(
        min_row=1, max_row=header_depth, max_col=1 + len(columns) * len(spec.measures)
    ):
        for cell in row:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4B123F")
    for index, node in enumerate(rows, start=header_depth + 1):
        label = text(index, 1, node["label"])
        label.alignment = Alignment(indent=max(0, node["depth"] - 1))
        subtotal = node["depth"] < len(spec.rows)
        label.font = Font(bold=subtotal)
        for column_index, column in enumerate(columns):
            values = cells.get((node["id"], column["id"]), [None] * len(spec.measures))
            for measure_index, measure in enumerate(spec.measures):
                target = sheet.cell(index, 2 + column_index * len(spec.measures) + measure_index)
                value = values[measure_index]
                if value is not None:
                    number = Decimal(value)
                    if len(number.as_tuple().digits) > 15:
                        # Preserve displayed text and raw SQL precision beyond Excel's limit.
                        with localcontext() as context:
                            context.prec = 300
                            display = number * (100 if measure.format == "percent" else 1)
                            display = display.quantize(
                                Decimal(1).scaleb(-measure.decimal_places),
                                rounding=ROUND_HALF_UP,
                            )
                            if not display:
                                display = abs(display)
                            formatted = f"{display:,.{measure.decimal_places}f}"
                        text(
                            index,
                            target.column,
                            formatted + ("%" if measure.format == "percent" else ""),
                        )
                        target.comment = Comment(f"Exact SQL value: {value}", "Data Agent")
                    else:
                        target.value = number
                    target.number_format = (
                        "#,##0"
                        + ("." + "0" * measure.decimal_places if measure.decimal_places else "")
                        + ("%" if measure.format == "percent" else "")
                    )
                target.font = Font(bold=subtotal or column["depth"] < len(spec.columns))
                target.alignment = Alignment(horizontal="right")
                sheet.column_dimensions[get_column_letter(target.column)].width = 18
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
