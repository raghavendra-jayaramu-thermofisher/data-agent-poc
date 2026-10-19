from __future__ import annotations

import re
from io import BytesIO
from typing import Annotated

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, ConfigDict, Field, model_validator

CellValue = Annotated[str, Field(max_length=32767)] | int | float | bool | None


class ExcelTableExport(BaseModel):
    model_config = ConfigDict(strict=True, allow_inf_nan=False)

    title: str = Field(default="Table", max_length=240)
    headers: list[Annotated[str, Field(max_length=512)]] = Field(min_length=1, max_length=256)
    rows: list[list[CellValue]] = Field(max_length=100000)
    truncated: bool = False

    @model_validator(mode="after")
    def validate_dimensions(self) -> ExcelTableExport:
        if any(len(row) != len(self.headers) for row in self.rows):
            raise ValueError("Every row must have one value for each column.")
        if len(self.headers) * len(self.rows) > 2_000_000:
            raise ValueError("This table is too large to export in one workbook.")
        return self


def export_filename(title: str) -> str:
    name = re.sub(r"[^A-Za-z0-9._ -]+", "_", title).strip(" ._")[:120]
    return f"{name or 'Table'}.xlsx"


def build_excel(table: ExcelTableExport) -> bytes:
    workbook = Workbook(write_only=True)
    workbook.properties.title = ILLEGAL_CHARACTERS_RE.sub("", table.title)
    title = re.sub(r"[\\/*?:\[\]\x00-\x1f]", "_", table.title).strip(" '")[:31]
    sheet = workbook.create_sheet(title or "Table")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(table.headers))}{len(table.rows) + 1}"
    for index, header in enumerate(table.headers):
        width = max([len(header), *(len(str(row[index] or "")) for row in table.rows[:200])])
        sheet.column_dimensions[get_column_letter(index + 1)].width = min(60, max(12, width + 2))

    def cell(value: CellValue, *, header: bool = False):
        if isinstance(value, int) and not isinstance(value, bool) and abs(value) >= 10**15:
            value = str(value)  # Excel numbers cannot retain identifiers longer than 15 digits.
        if isinstance(value, str):
            value = ILLEGAL_CHARACTERS_RE.sub("", value)
        item = WriteOnlyCell(sheet, value=value)
        if isinstance(value, str):
            item.data_type = "s"  # Preserve literal text such as =SUM(...), never create formulas.
        if header:
            item.font = Font(bold=True, color="FFFFFF")
            item.fill = PatternFill("solid", fgColor="5B5FC7")
        return item

    sheet.append([cell(header, header=True) for header in table.headers])
    for row in table.rows:
        sheet.append([cell(value) for value in row])
    if table.truncated:
        workbook.properties.description = (
            f"Includes all {len(table.rows)} available rows. The source result was truncated."
        )
        sheet.append([])
        sheet.append(
            [cell("The source result was truncated; this export contains available rows only.")]
        )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
