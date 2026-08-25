from __future__ import annotations

import asyncio
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Literal

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.core.config import Settings
from app.data_sources.errors import DataSourceOperationError

LogicalType = Literal[
    "unknown", "boolean", "integer", "double", "date", "timestamp", "time", "text"
]


@dataclass(frozen=True, slots=True)
class SheetMaterialization:
    original_sheet_name: str
    table_name: str | None
    row_count: int
    column_count: int
    skipped: bool

    def as_dict(self) -> dict[str, object]:
        return {
            "original_sheet_name": self.original_sheet_name,
            "table_name": self.table_name,
            "row_count": self.row_count,
            "column_count": self.column_count,
            "skipped": self.skipped,
        }


@dataclass(frozen=True, slots=True)
class MaterializationResult:
    sheets: tuple[SheetMaterialization, ...]


_IDENTIFIER_RE = re.compile(r"[^a-z0-9_]+")


def normalize_identifier(value: object, *, prefix: str, fallback_index: int | None = None) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode()
    text = _IDENTIFIER_RE.sub("_", text.lower()).strip("_")
    if not text:
        suffix = str(fallback_index or 1)
        text = f"{prefix}_{suffix}"
    if text[0].isdigit():
        text = f"{prefix}_{text}"
    return text[:120]


def unique_identifiers(values: list[object], *, prefix: str) -> list[str]:
    counts: dict[str, int] = {}
    result: list[str] = []
    for index, value in enumerate(values, start=1):
        base = normalize_identifier(value, prefix=prefix, fallback_index=index)
        count = counts.get(base, 0) + 1
        counts[base] = count
        result.append(base if count == 1 else f"{base}_{count}")
    return result


def classify_value(value: object) -> LogicalType:
    if value is None:
        return "unknown"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int) and not isinstance(value, bool):
        return "integer"
    if isinstance(value, (float, Decimal)):
        return "double"
    if isinstance(value, datetime):
        return "timestamp"
    if isinstance(value, date):
        return "date"
    if isinstance(value, time):
        return "time"
    return "text"


def merge_types(current: LogicalType, incoming: LogicalType) -> LogicalType:
    if incoming == "unknown":
        return current
    if current == "unknown":
        return incoming
    if current == incoming:
        return current
    if {current, incoming} <= {"integer", "double"}:
        return "double"
    if {current, incoming} <= {"date", "timestamp"}:
        return "timestamp"
    # Any other disagreement falls back to text so values are preserved rather than coerced away.
    return "text"


def duckdb_type(logical_type: LogicalType) -> str:
    return {
        "unknown": "VARCHAR",
        "boolean": "BOOLEAN",
        "integer": "BIGINT",
        "double": "DOUBLE",
        "date": "DATE",
        "timestamp": "TIMESTAMP",
        "time": "TIME",
        "text": "VARCHAR",
    }[logical_type]


def convert_value(value: object, logical_type: LogicalType) -> object:
    if value is None:
        return None
    if logical_type == "text":
        return str(value)
    if logical_type == "double" and isinstance(value, Decimal):
        return float(value)
    if logical_type == "timestamp" and isinstance(value, date) and not isinstance(value, datetime):
        return datetime.combine(value, time.min)
    return value


def quote_identifier(value: str) -> str:
    return '"' + value.replace('"', '""') + '"'


class ExcelMaterializer:
    def __init__(self, settings: Settings) -> None:
        self.batch_size = settings.excel_batch_size

    @staticmethod
    def _duckdb():  # type: ignore[no-untyped-def]
        try:
            import duckdb
        except ImportError as exc:
            raise DataSourceOperationError(
                "DATASOURCE_DRIVER_UNAVAILABLE",
                "DuckDB support is not installed on this server.",
            ) from exc
        return duckdb

    async def list_sheets(self, source_path: Path) -> list[str]:
        return await asyncio.to_thread(self._list_sheets_sync, source_path)

    @staticmethod
    def _list_sheets_sync(source_path: Path) -> list[str]:
        try:
            workbook = load_workbook(source_path, read_only=True, data_only=True)
        except (InvalidFileException, OSError, ValueError, KeyError) as exc:
            raise DataSourceOperationError(
                "EXCEL_WORKBOOK_INVALID", "The uploaded workbook could not be read."
            ) from exc
        try:
            return list(workbook.sheetnames)
        finally:
            workbook.close()

    async def initialize_database(self, database_path: Path) -> None:
        await asyncio.to_thread(self._initialize_database_sync, database_path)

    def _initialize_database_sync(self, database_path: Path) -> None:
        database_path.parent.mkdir(parents=True, exist_ok=True)
        database_path.unlink(missing_ok=True)
        connection = self._duckdb().connect(str(database_path))
        try:
            connection.execute("CHECKPOINT")
        finally:
            connection.close()

    async def materialize_sheet(
        self,
        *,
        source_path: Path,
        database_path: Path,
        sheet_name: str,
        table_name: str,
    ) -> SheetMaterialization:
        return await asyncio.to_thread(
            self._materialize_sheet_sync,
            source_path,
            database_path,
            sheet_name,
            table_name,
        )

    def _materialize_sheet_sync(
        self, source_path: Path, database_path: Path, sheet_name: str, table_name: str
    ) -> SheetMaterialization:
        try:
            workbook = load_workbook(source_path, read_only=True, data_only=True)
        except (InvalidFileException, OSError, ValueError, KeyError) as exc:
            raise DataSourceOperationError(
                "EXCEL_WORKBOOK_INVALID", "The uploaded workbook could not be read."
            ) from exc
        try:
            worksheet = workbook[sheet_name]
            max_column = int(worksheet.max_column or 0)
            max_row = int(worksheet.max_row or 0)
            if max_column <= 0 or max_row <= 0:
                return SheetMaterialization(sheet_name, None, 0, 0, True)

            first_row = next(
                worksheet.iter_rows(min_row=1, max_row=1, max_col=max_column, values_only=True),
                (),
            )
            if not first_row or all(value is None for value in first_row) and max_row == 1:
                return SheetMaterialization(sheet_name, None, 0, 0, True)
            columns = unique_identifiers(list(first_row), prefix="column")
            logical_types: list[LogicalType] = ["unknown"] * len(columns)
            row_count = 0
            for row in worksheet.iter_rows(
                min_row=2, max_row=max_row, max_col=len(columns), values_only=True
            ):
                if all(value is None for value in row):
                    continue
                row_count += 1
                for index, value in enumerate(row):
                    logical_types[index] = merge_types(logical_types[index], classify_value(value))
        finally:
            workbook.close()

        connection = self._duckdb().connect(str(database_path))
        try:
            columns_sql = ", ".join(
                f"{quote_identifier(name)} {duckdb_type(logical_type)}"
                for name, logical_type in zip(columns, logical_types, strict=True)
            )
            connection.execute(f"DROP TABLE IF EXISTS {quote_identifier(table_name)}")
            connection.execute(f"CREATE TABLE {quote_identifier(table_name)} ({columns_sql})")
            if row_count == 0:
                return SheetMaterialization(sheet_name, table_name, 0, len(columns), False)

            try:
                workbook = load_workbook(source_path, read_only=True, data_only=True)
            except (InvalidFileException, OSError, ValueError, KeyError) as exc:
                raise DataSourceOperationError(
                    "EXCEL_WORKBOOK_INVALID", "The uploaded workbook could not be read."
                ) from exc
            try:
                worksheet = workbook[sheet_name]
                placeholders = ", ".join(["?"] * len(columns))
                insert_sql = (
                    f"INSERT INTO {quote_identifier(table_name)} VALUES ({placeholders})"
                )
                batch: list[tuple[object, ...]] = []
                for row in worksheet.iter_rows(
                    min_row=2,
                    max_row=worksheet.max_row,
                    max_col=len(columns),
                    values_only=True,
                ):
                    if all(value is None for value in row):
                        continue
                    batch.append(
                        tuple(
                            convert_value(value, logical_types[index])
                            for index, value in enumerate(row)
                        )
                    )
                    if len(batch) >= self.batch_size:
                        connection.executemany(insert_sql, batch)
                        batch.clear()
                if batch:
                    connection.executemany(insert_sql, batch)
            finally:
                workbook.close()
            return SheetMaterialization(sheet_name, table_name, row_count, len(columns), False)
        except DataSourceOperationError:
            raise
        except Exception as exc:
            raise DataSourceOperationError(
                "EXCEL_MATERIALIZATION_FAILED",
                "The workbook could not be converted into an analytical database.",
            ) from exc
        finally:
            connection.close()

    async def validate_database(self, database_path: Path) -> None:
        await asyncio.to_thread(self._validate_database_sync, database_path)

    def _validate_database_sync(self, database_path: Path) -> None:
        connection = self._duckdb().connect(str(database_path), read_only=True)
        try:
            count = connection.execute(
                """
                SELECT COUNT(*)
                FROM information_schema.tables
                WHERE table_schema NOT IN ('information_schema', 'pg_catalog')
                  AND table_type = 'BASE TABLE'
                """
            ).fetchone()[0]
            if int(count) < 1:
                raise DataSourceOperationError(
                    "EXCEL_NO_USABLE_SHEETS", "The workbook contains no usable worksheets."
                )
            connection.execute("SELECT 1").fetchone()
        except DataSourceOperationError:
            raise
        except Exception as exc:
            raise DataSourceOperationError(
                "EXCEL_MATERIALIZATION_INVALID",
                "The converted analytical database could not be validated.",
            ) from exc
        finally:
            connection.close()
